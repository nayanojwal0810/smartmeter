"""REFIT Dataset Forensic Quality Auditor.

Streaming, low-memory data validation and audit engine for cleaned REFIT CSV files.
Performs schema verification, temporal monotonicity and gap distribution analysis,
power value statistics, issue flag frequencies, and metadata-channel mapping.
"""

from __future__ import annotations

import csv
import json
import logging
import math
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

logger = logging.getLogger(__name__)

EXPECTED_HEADER = [
    "Time",
    "Unix",
    "Aggregate",
    "Appliance1",
    "Appliance2",
    "Appliance3",
    "Appliance4",
    "Appliance5",
    "Appliance6",
    "Appliance7",
    "Appliance8",
    "Appliance9",
    "Issues",
]


@dataclass
class ColumnStats:
    """Streaming numerical summary for a single power channel."""

    valid_count: int = 0
    sum_val: float = 0.0
    min_val: float = float("inf")
    max_val: float = float("-inf")
    zero_count: int = 0
    negative_count: int = 0
    above_4000_count: int = 0
    nan_count: int = 0
    non_numeric_count: int = 0

    @property
    def mean_val(self) -> Optional[float]:
        """Compute arithmetic mean of valid observations."""
        if self.valid_count == 0:
            return None
        return self.sum_val / self.valid_count

    def update(self, val_str: str) -> None:
        """Update running statistics with a raw string observation."""
        val_clean = val_str.strip()
        if not val_clean:
            self.non_numeric_count += 1
            return

        if val_clean.lower() in ("nan", "null", "none"):
            self.nan_count += 1
            return

        try:
            val = float(val_clean)
        except ValueError:
            self.non_numeric_count += 1
            return

        if math.isnan(val):
            self.nan_count += 1
            return
        if math.isinf(val):
            self.non_numeric_count += 1
            return

        self.valid_count += 1
        self.sum_val += val
        if val < self.min_val:
            self.min_val = val
        if val > self.max_val:
            self.max_val = val

        if val == 0.0:
            self.zero_count += 1
        elif val < 0.0:
            self.negative_count += 1

        if val > 4000.0:
            self.above_4000_count += 1

    def to_dict(self) -> Dict[str, Any]:
        """Convert statistics to serializable dictionary."""
        return {
            "valid_count": self.valid_count,
            "mean_val": round(self.mean_val, 3) if self.mean_val is not None else None,
            "min_val": self.min_val if self.valid_count > 0 else None,
            "max_val": self.max_val if self.valid_count > 0 else None,
            "zero_count": self.zero_count,
            "zero_pct": round(100.0 * self.zero_count / self.valid_count, 2) if self.valid_count > 0 else 0.0,
            "negative_count": self.negative_count,
            "above_4000_count": self.above_4000_count,
            "nan_count": self.nan_count,
            "non_numeric_count": self.non_numeric_count,
        }


@dataclass
class GapStats:
    """Sampling interval and time gap distribution statistics."""

    total_intervals: int = 0
    le_8s: int = 0  # <= 8s
    between_9_15s: int = 0  # 9s to 15s
    gt_15s: int = 0  # > 15s
    gt_30s: int = 0  # > 30s
    gt_60s: int = 0  # > 60s
    gt_120s: int = 0  # > 120s
    zero_or_negative_gaps: int = 0
    max_gap_seconds: float = 0.0
    min_gap_seconds: float = float("inf")

    def update(self, gap_seconds: float) -> None:
        """Record a time step gap between consecutive records."""
        self.total_intervals += 1
        if gap_seconds < self.min_gap_seconds:
            self.min_gap_seconds = gap_seconds
        if gap_seconds > self.max_gap_seconds:
            self.max_gap_seconds = gap_seconds

        if gap_seconds <= 0:
            self.zero_or_negative_gaps += 1
            return

        if gap_seconds <= 8.0:
            self.le_8s += 1
        elif gap_seconds <= 15.0:
            self.between_9_15s += 1
        else:
            self.gt_15s += 1

        if gap_seconds > 30.0:
            self.gt_30s += 1
        if gap_seconds > 60.0:
            self.gt_60s += 1
        if gap_seconds > 120.0:
            self.gt_120s += 1

    def to_dict(self) -> Dict[str, Any]:
        """Convert gap distribution to serializable dictionary."""
        total = max(1, self.total_intervals)
        return {
            "total_intervals": self.total_intervals,
            "min_gap_seconds": self.min_gap_seconds if self.total_intervals > 0 else None,
            "max_gap_seconds": self.max_gap_seconds,
            "le_8s_count": self.le_8s,
            "le_8s_pct": round(100.0 * self.le_8s / total, 2),
            "between_9_15s_count": self.between_9_15s,
            "between_9_15s_pct": round(100.0 * self.between_9_15s / total, 2),
            "gt_15s_count": self.gt_15s,
            "gt_15s_pct": round(100.0 * self.gt_15s / total, 2),
            "gt_30s_count": self.gt_30s,
            "gt_30s_pct": round(100.0 * self.gt_30s / total, 2),
            "gt_60s_count": self.gt_60s,
            "gt_60s_pct": round(100.0 * self.gt_60s / total, 2),
            "gt_120s_count": self.gt_120s,
            "gt_120s_pct": round(100.0 * self.gt_120s / total, 2),
            "zero_or_negative_gaps": self.zero_or_negative_gaps,
        }


@dataclass
class HouseholdAuditResult:
    """Comprehensive audit metrics for an individual household CSV file."""

    household_id: int
    filename: str
    file_size_bytes: int
    total_rows: int = 0
    malformed_rows: int = 0
    header_valid: bool = False
    actual_header: List[str] = field(default_factory=list)
    min_datetime: str = ""
    max_datetime: str = ""
    min_unix: Optional[int] = None
    max_unix: Optional[int] = None
    is_monotonic: bool = True
    duplicate_timestamps: int = 0
    gap_stats: GapStats = field(default_factory=GapStats)
    aggregate_stats: ColumnStats = field(default_factory=ColumnStats)
    appliance_stats: Dict[str, ColumnStats] = field(default_factory=dict)
    issues_total: int = 0
    issues_1_count: int = 0
    issues_0_count: int = 0
    issues_invalid_count: int = 0
    appliance_mapping: Dict[str, Dict[str, str]] = field(default_factory=dict)

    def to_dict(self) -> Dict[str, Any]:
        """Convert household audit result to dictionary."""
        issues_pct = round(100.0 * self.issues_1_count / max(1, self.total_rows), 4)
        return {
            "household_id": self.household_id,
            "filename": self.filename,
            "file_size_bytes": self.file_size_bytes,
            "total_rows": self.total_rows,
            "malformed_rows": self.malformed_rows,
            "header_valid": self.header_valid,
            "min_datetime": self.min_datetime,
            "max_datetime": self.max_datetime,
            "min_unix": self.min_unix,
            "max_unix": self.max_unix,
            "is_monotonic": self.is_monotonic,
            "duplicate_timestamps": self.duplicate_timestamps,
            "gaps": self.gap_stats.to_dict(),
            "aggregate": self.aggregate_stats.to_dict(),
            "appliances": {k: v.to_dict() for k, v in self.appliance_stats.items()},
            "issues": {
                "total": self.issues_total,
                "count_1": self.issues_1_count,
                "count_0": self.issues_0_count,
                "count_invalid": self.issues_invalid_count,
                "pct_1": issues_pct,
            },
            "appliance_mapping": self.appliance_mapping,
        }


def load_refit_metadata(metadata_path: Path) -> Dict[int, Dict[str, Any]]:
    """Parse REFIT metadata workbook and extract appliance mappings per house.

    Args:
        metadata_path: Path to MetaData_Tables.xlsx

    Returns:
        Dictionary mapping household ID (int) to house demographic info and
        IAM channel mappings (IAM 1..9 -> name, make, model).
    """
    try:
        import openpyxl
    except ImportError as e:
        logger.error("openpyxl is required to read REFIT metadata tables.")
        raise RuntimeError("openpyxl must be installed to parse MetaData_Tables.xlsx") from e

    wb = openpyxl.load_workbook(str(metadata_path), data_only=True)
    metadata: Dict[int, Dict[str, Any]] = {}

    # 1. Parse demographics from Sheet1
    if "Sheet1" in wb.sheetnames:
        sheet1 = wb["Sheet1"]
        for row in sheet1.iter_rows(values_only=True):
            if not row or not row[0]:
                continue
            first_val = str(row[0]).strip()
            if first_val.isdigit():
                h_id = int(first_val)
                metadata[h_id] = {
                    "occupancy": str(row[1]).strip() if len(row) > 1 and row[1] is not None else "",
                    "construction_year": str(row[2]).strip() if len(row) > 2 and row[2] is not None else "",
                    "appliances_owned_count": int(row[3]) if len(row) > 3 and isinstance(row[3], (int, float)) else 0,
                    "house_type": str(row[4]).strip() if len(row) > 4 and row[4] is not None else "",
                    "house_size": str(row[5]).strip() if len(row) > 5 and row[5] is not None else "",
                    "iam_channels": {},
                }

    # 2. Parse individual house sheets
    for sheet_name in wb.sheetnames:
        if not sheet_name.startswith("House "):
            continue
        try:
            h_id = int(sheet_name.replace("House ", "").strip())
        except ValueError:
            continue

        if h_id not in metadata:
            metadata[h_id] = {
                "occupancy": "",
                "construction_year": "",
                "appliances_owned_count": 0,
                "house_type": "",
                "house_size": "",
                "iam_channels": {},
            }

        sheet = wb[sheet_name]
        for row in sheet.iter_rows(values_only=True):
            if not row or row[0] is None:
                continue
            try:
                iam_idx = int(row[0])
            except ValueError:
                continue

            if iam_idx == 0:
                continue  # Aggregate channel

            col_key = f"Appliance{iam_idx}"
            app_name = str(row[1]).strip() if len(row) > 1 and row[1] is not None else "Unknown"
            make = str(row[2]).strip() if len(row) > 2 and row[2] is not None else ""
            model = str(row[3]).strip() if len(row) > 3 and row[3] is not None else ""

            metadata[h_id]["iam_channels"][col_key] = {
                "iam_index": iam_idx,
                "name": app_name,
                "make": make,
                "model": model,
            }

    return metadata


class REFITAuditor:
    """Forensic auditor for streaming REFIT time series data."""

    def __init__(self, metadata_path: Optional[Path] = None) -> None:
        self.metadata_path = metadata_path
        self.metadata: Dict[int, Dict[str, Any]] = {}
        if metadata_path and metadata_path.exists():
            self.metadata = load_refit_metadata(metadata_path)

    def audit_household_file(
        self,
        file_path: Path,
        household_id: Optional[int] = None,
    ) -> HouseholdAuditResult:
        """Stream through a single household CSV and compute audit metrics.

        Memory complexity: O(1). Does not load file contents into RAM.

        Args:
            file_path: Path to CLEAN_HouseN.csv
            household_id: Optional integer ID. Inferred from filename if omitted.

        Returns:
            Populated HouseholdAuditResult object.
        """
        filename = file_path.name
        file_size = file_path.stat().st_size

        if household_id is None:
            # Parse from filename e.g. CLEAN_House3.csv -> 3
            digits = "".join(c for c in filename if c.isdigit())
            household_id = int(digits) if digits else 0

        iam_map = {}
        if household_id in self.metadata:
            iam_map = self.metadata[household_id].get("iam_channels", {})

        result = HouseholdAuditResult(
            household_id=household_id,
            filename=filename,
            file_size_bytes=file_size,
            appliance_mapping=iam_map,
        )

        for i in range(1, 10):
            result.appliance_stats[f"Appliance{i}"] = ColumnStats()

        prev_unix: Optional[int] = None

        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            reader = csv.reader(f)
            try:
                header = next(reader)
            except StopIteration:
                return result

            result.actual_header = header
            result.header_valid = header == EXPECTED_HEADER

            for row in reader:
                result.total_rows += 1
                if len(row) != 13:
                    result.malformed_rows += 1
                    continue

                # 1. Timestamps
                dt_str, unix_str = row[0].strip(), row[1].strip()
                if not result.min_datetime:
                    result.min_datetime = dt_str
                result.max_datetime = dt_str

                try:
                    unix_val = int(unix_str)
                except ValueError:
                    result.malformed_rows += 1
                    continue

                if result.min_unix is None or unix_val < result.min_unix:
                    result.min_unix = unix_val
                if result.max_unix is None or unix_val > result.max_unix:
                    result.max_unix = unix_val

                if prev_unix is not None:
                    gap = float(unix_val - prev_unix)
                    if gap <= 0:
                        result.is_monotonic = False
                        if gap == 0:
                            result.duplicate_timestamps += 1
                    result.gap_stats.update(gap)
                prev_unix = unix_val

                # 2. Aggregate power
                result.aggregate_stats.update(row[2])

                # 3. Appliance channels 1 to 9
                for i in range(1, 10):
                    result.appliance_stats[f"Appliance{i}"].update(row[2 + i])

                # 4. Issues column
                issue_val = row[12].strip()
                result.issues_total += 1
                if issue_val == "1":
                    result.issues_1_count += 1
                elif issue_val == "0":
                    result.issues_0_count += 1
                else:
                    result.issues_invalid_count += 1

        return result

    def audit_all_households(
        self,
        data_dir: Path,
    ) -> List[HouseholdAuditResult]:
        """Audit all CLEAN_House*.csv files in directory in sorted order."""
        files = sorted(
            data_dir.glob("CLEAN_House*.csv"),
            key=lambda p: int("".join(filter(str.isdigit, p.name)) or 0),
        )
        results: List[HouseholdAuditResult] = []
        for file_path in files:
            logger.info("Auditing file: %s", file_path.name)
            res = self.audit_household_file(file_path)
            results.append(res)
        return results

    def generate_summary_report(
        self,
        results: List[HouseholdAuditResult],
    ) -> Dict[str, Any]:
        """Produce consolidated cross-household summary dictionary."""
        total_files = len(results)
        total_bytes = sum(r.file_size_bytes for r in results)
        total_rows = sum(r.total_rows for r in results)
        total_issues_1 = sum(r.issues_1_count for r in results)
        all_headers_valid = all(r.header_valid for r in results)
        all_monotonic = all(r.is_monotonic for r in results)

        # Cross-household gap totals
        total_intervals = sum(r.gap_stats.total_intervals for r in results)
        total_le_8s = sum(r.gap_stats.le_8s for r in results)
        total_9_15s = sum(r.gap_stats.between_9_15s for r in results)
        total_gt_15s = sum(r.gap_stats.gt_15s for r in results)
        total_gt_30s = sum(r.gap_stats.gt_30s for r in results)
        total_gt_60s = sum(r.gap_stats.gt_60s for r in results)
        total_gt_120s = sum(r.gap_stats.gt_120s for r in results)
        max_overall_gap = max((r.gap_stats.max_gap_seconds for r in results), default=0.0)

        # Candidate appliance occurrences across audited households
        appliance_occurrences: Dict[str, List[int]] = {}
        for r in results:
            for col_name, app_info in r.appliance_mapping.items():
                name = app_info.get("name", "Unknown")
                if name not in appliance_occurrences:
                    appliance_occurrences[name] = []
                appliance_occurrences[name].append(r.household_id)

        return {
            "overview": {
                "total_households": total_files,
                "total_size_bytes": total_bytes,
                "total_size_gb": round(total_bytes / (1024**3), 3),
                "total_rows": total_rows,
                "total_issues_1_rows": total_issues_1,
                "overall_issues_pct": round(100.0 * total_issues_1 / max(1, total_rows), 4),
                "all_headers_valid": all_headers_valid,
                "all_monotonic": all_monotonic,
            },
            "aggregate_sampling_gaps": {
                "total_intervals": total_intervals,
                "le_8s_pct": round(100.0 * total_le_8s / max(1, total_intervals), 2),
                "between_9_15s_pct": round(100.0 * total_9_15s / max(1, total_intervals), 2),
                "gt_15s_pct": round(100.0 * total_gt_15s / max(1, total_intervals), 2),
                "gt_30s_pct": round(100.0 * total_gt_30s / max(1, total_intervals), 2),
                "gt_60s_pct": round(100.0 * total_gt_60s / max(1, total_intervals), 2),
                "gt_120s_pct": round(100.0 * total_gt_120s / max(1, total_intervals), 2),
                "max_gap_seconds": max_overall_gap,
            },
            "appliance_presence_summary": {
                k: {"household_count": len(v), "households": sorted(v)}
                for k, v in sorted(appliance_occurrences.items(), key=lambda item: len(item[1]), reverse=True)
            },
            "households": [r.to_dict() for r in results],
        }

    def generate_markdown_report(self, summary: Dict[str, Any]) -> str:
        """Render a formatted GitHub-style Markdown audit report."""
        ov = summary["overview"]
        gaps = summary["aggregate_sampling_gaps"]
        apps = summary["appliance_presence_summary"]
        households = summary["households"]

        lines = [
            "# REFIT Dataset Forensic Data Quality Audit Report",
            "",
            "## 1. Executive Summary",
            "",
            "| Metric | Value |",
            "|---|---|",
            f"| **Total Households Audited** | {ov['total_households']} |",
            f"| **Total Uncompressed Volume** | {ov['total_size_gb']} GB ({ov['total_size_bytes']:,} bytes) |",
            f"| **Total Recorded Rows** | {ov['total_rows']:,} |",
            f"| **Total Rows with Issues == 1** | {ov['total_issues_1_rows']:,} ({ov['overall_issues_pct']}%) |",
            f"| **Schema Conformance** | {'100% Valid (All files match exact 13-column schema)' if ov['all_headers_valid'] else 'Schema mismatches detected'} |",
            f"| **Temporal Monotonicity** | {'Monotonic (no timestamp reversals)' if ov['all_monotonic'] else 'Reversals or non-monotonic timestamps detected'} |",
            "",
            "## 2. Sampling Interval & Gap Distribution",
            "",
            "| Interval Category | Count / Percentage | Description |",
            "|---|---|---|",
            f"| **<= 8 seconds** | {gaps['le_8s_pct']}% | Nominal target sampling rate |",
            f"| **9–15 seconds** | {gaps['between_9_15s_pct']}% | Minor polling / load-change delay |",
            f"| **> 15 seconds** | {gaps['gt_15s_pct']}% | Irregular transmission gap |",
            f"| **> 30 seconds** | {gaps['gt_30s_pct']}% | Medium gap |",
            f"| **> 60 seconds** | {gaps['gt_60s_pct']}% | Extended gap |",
            f"| **> 120 seconds** | {gaps['gt_120s_pct']}% | Substantial missing outage |",
            f"| **Maximum Observed Gap** | {gaps['max_gap_seconds']:,.1f} s ({gaps['max_gap_seconds']/86400:.2f} days) | Maximum sensor/network outage |",
            "",
            "## 3. Appliance Channel Distribution Across Households",
            "",
            "| Appliance Name | Monitored Household Count | Household IDs |",
            "|---|---|---|",
        ]

        for app_name, info in apps.items():
            h_str = ", ".join(f"H{h}" for h in info["households"])
            lines.append(f"| **{app_name}** | {info['household_count']} | {h_str} |")

        lines.extend([
            "",
            "## 4. Household-by-Household Summary Table",
            "",
            "| House | Rows | Start Date | End Date | Aggregate Mean (W) | Aggregate Max (W) | Issues % | Monotonic | Max Gap (h) |",
            "|---|---|---|---|---|---|---|---|---|",
        ])

        for h in households:
            h_id = h["household_id"]
            rows = f"{h['total_rows']:,}"
            start = h["min_datetime"].split()[0] if h["min_datetime"] else "N/A"
            end = h["max_datetime"].split()[0] if h["max_datetime"] else "N/A"
            agg_mean = f"{h['aggregate']['mean_val']:.1f}" if h['aggregate']['mean_val'] else "N/A"
            agg_max = f"{h['aggregate']['max_val']:.0f}" if h['aggregate']['max_val'] else "N/A"
            issues_pct = f"{h['issues']['pct_1']:.2f}%"
            mono = "Yes" if h["is_monotonic"] else "No"
            max_gap_h = f"{h['gaps']['max_gap_seconds']/3600:.1f}h"
            lines.append(
                f"| House {h_id} | {rows} | {start} | {end} | {agg_mean} | {agg_max} | {issues_pct} | {mono} | {max_gap_h} |"
            )

        return "\n".join(lines)
